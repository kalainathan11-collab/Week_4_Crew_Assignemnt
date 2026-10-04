# First time setup: pip install python-docx

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def generate_orangehrm_brd():
    doc = Document()

    # Color Palette Definitions
    NAVY = RGBColor(0x1B, 0x36, 0x5D)
    GREY = RGBColor(0x55, 0x55, 0x55)

    # Document Header
    title = doc.add_heading('Business Requirements Document (BRD)', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.color.rgb = NAVY

    subtitle = doc.add_paragraph('UI Application: OrangeHRM Open Source Management Portal')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.italic = True
    subtitle.runs[0].font.color.rgb = GREY

    doc.add_paragraph('Target URL: https://opensource-demo.orangehrmlive.com/\nVersion: 1.0 | Date: October 2026 | Author: Lead Business Analyst\n')

    # 1. Executive Summary
    h1 = doc.add_heading('1. Executive Summary', level=1)
    h1.runs[0].font.color.rgb = NAVY
    doc.add_paragraph(
        "The OrangeHRM Web Application is an HR Management System (HRMS) designed to centralize core HR "
        "workflows, employee records, attendance tracking, and performance reviews into a responsive single-page application (SPA)."
    )

    # 2. User Roles Matrix Table
    h2 = doc.add_heading('2. Target User Roles & Access Matrix', level=1)
    h2.runs[0].font.color.rgb = NAVY

    roles_table = doc.add_table(rows=4, cols=3)
    roles_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    roles_table.style = 'Table Grid'

    r_headers = ['Role', 'Access Level', 'Primary UI Scope']
    for idx, text in enumerate(r_headers):
        cell = roles_table.cell(0, idx)
        cell.text = text
        cell.paragraphs[0].runs[0].font.bold = True

    r_data = [
        ('System Admin', 'Full Read/Write', 'System settings, user management, custom reports'),
        ('HR Manager', 'Module Admin', 'Employee directory, leave approvals, recruitment'),
        ('Employee', 'Self-Service', 'Personal profile, leave application, timesheets')
    ]

    for row_idx, data in enumerate(r_data, start=1):
        for col_idx, text in enumerate(data):
            roles_table.cell(row_idx, col_idx).text = text

    doc.add_paragraph('')

    # 3. Functional Requirements
    h3 = doc.add_heading('3. UI Functional Requirements', level=1)
    h3.runs[0].font.color.rgb = NAVY

    req_table = doc.add_table(rows=7, cols=4)
    req_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    req_table.style = 'Table Grid'

    f_headers = ['Req ID', 'Module', 'Priority', 'Description']
    for idx, text in enumerate(f_headers):
        cell = req_table.cell(0, idx)
        cell.text = text
        cell.paragraphs[0].runs[0].font.bold = True

    f_data = [
        ('FR-AUTH-01', 'Authentication', 'Critical', 'Login card with Username/Password inputs and floating labels.'),
        ('FR-AUTH-02', 'Authentication', 'Critical', 'Login button triggering authentication with loading state.'),
        ('FR-DASH-01', 'Dashboard', 'High', 'Quick Launch grid shortcuts for Leave and Timesheets.'),
        ('FR-DASH-02', 'Dashboard', 'High', 'Donut chart rendering employee distribution by unit.'),
        ('FR-PIM-01', 'PIM Module', 'Critical', 'Search filter accordion for Employee Name, ID, and Job Title.'),
        ('FR-PIM-02', 'PIM Module', 'Critical', 'Paginated grid showing employee ID, Name, Job Title, and Actions.')
    ]

    for row_idx, data in enumerate(f_data, start=1):
        for col_idx, text in enumerate(data):
            req_table.cell(row_idx, col_idx).text = text

    doc.add_paragraph('')

    # 4. Non-Functional Requirements
    h4 = doc.add_heading('4. Non-Functional Requirements', level=1)
    h4.runs[0].font.color.rgb = NAVY

    doc.add_paragraph('• Accessibility: Compliant with WCAG 2.1 Level AA contrast and keyboard standards.')
    doc.add_paragraph('• Performance: Time-to-Interactive (TTI) under 1.8 seconds on corporate broadband.')
    doc.add_paragraph('• Session Timeout: Automatic logout after 15 minutes of user inactivity.')

    # Save output
    output_filename = 'OrangeHRM_BRD_Document.docx'
    doc.save(output_filename)
    print(f"File successfully created: '{output_filename}'")

if __name__ == '__main__':
    generate_orangehrm_brd()